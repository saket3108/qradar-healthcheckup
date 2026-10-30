"""Editable DOCX report (python-docx) with native tables, headings (navigable TOC) and embedded charts."""
from __future__ import annotations

import io
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.reports import charts

NAVY, BLUE, GREY = RGBColor(0x0F, 0x2A, 0x4A), RGBColor(0x0F, 0x62, 0xFE), RGBColor(0x6F, 0x77, 0x82)
SEV_HEX = {"critical": "DA1E28", "high": "FF832B", "medium": "D2A106", "low": "24A148", "info": "0F62FE"}


def _shade(cell, hex_fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _add_field(paragraph, instr: str):
    run = paragraph.add_run()
    for t, txt in (("begin", None), (None, instr), ("separate", None), ("end", None)):
        if t:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), t)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        run._r.append(el)


def _table(doc: Document, rows: list[list[Any]], widths_cm: list[float] | None = None, sev_col: int | None = None):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, r in enumerate(rows):
        for ci, v in enumerate(r):
            c = t.cell(ri, ci)
            c.text = ""
            p = c.paragraphs[0]
            run = p.add_run(str(v if v is not None else ""))
            run.font.size = Pt(8.5)
            if ri == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                _shade(c, "0F2A4A")
            else:
                if ri % 2 == 0:
                    _shade(c, "F2F4F8")
                if sev_col is not None and ci == sev_col and str(v).lower() in SEV_HEX:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(SEV_HEX[str(v).lower()])
            if widths_cm:
                c.width = Cm(widths_cm[ci])
    # repeat header row
    trPr = t.rows[0]._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader")
    h.set(qn("w:val"), "true")
    trPr.append(h)
    doc.add_paragraph()
    return t


def _img(doc: Document, png: bytes, width_cm: float = 16):
    doc.add_picture(io.BytesIO(png), width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def _p(doc: Document, text: str, bold: bool = False, italic: bool = False, size: float = 10, color: RGBColor | None = None, style: str | None = None):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    r.bold, r.italic, r.font.size = bold, italic, Pt(size)
    if color:
        r.font.color.rgb = color
    return p


def _bullets(doc: Document, items: list[str]):
    for i in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(i).font.size = Pt(9.5)


def _comment(doc: Document, text: str | None):
    if text:
        p = doc.add_paragraph()
        r = p.add_run("Consultant comment: ")
        r.bold = True
        r.font.size = Pt(9.5)
        p.add_run(text).font.size = Pt(9.5)
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "FFF8E1")
        pPr.append(shd)


def _gb(b):
    return f"{b / 1024**3:.0f} GB" if b else "-"


def build_docx(model: dict[str, Any]) -> bytes:
    a, cfg = model["assessment"], model["config"]
    secs, comments = cfg["sections"], model["comments"]
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10)
    for name, size, color in (("Heading 1", 20, NAVY), ("Heading 2", 14, BLUE), ("Heading 3", 11, NAVY)):
        hs = doc.styles[name]
        hs.font.name, hs.font.size, hs.font.color.rgb, hs.font.bold = "Calibri", Pt(size), color, True
    doc.core_properties.title, doc.core_properties.author = cfg["title"], a["consultant"] or "IBM Security"
    # header / footer
    hp = sec.header.paragraphs[0]
    hr = hp.add_run(f"{cfg['title']}  |  {a['customer']}  |  {cfg['classification']}")
    hr.font.size, hr.font.color.rgb = Pt(8), GREY
    fp = sec.footer.paragraphs[0]
    fr = fp.add_run(f"Prepared by {a['consultant'] or 'IBM Security'} - {a['generated']}    Page ")
    fr.font.size, fr.font.color.rgb = Pt(8), GREY
    _add_field(fp, "PAGE")

    def sc(key):
        _comment(doc, comments.get(f"section:{key}"))

    def pb():
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    if "cover" in secs:
        for _ in range(5):
            doc.add_paragraph()
        logo = cfg.get("logo_path")
        if logo and Path(logo).exists():
            try:
                doc.add_picture(logo, width=Cm(5))
            except Exception:  # noqa: BLE001
                pass
        p = doc.add_paragraph()
        r = p.add_run(cfg["title"])
        r.font.size, r.bold, r.font.color.rgb = Pt(32), True, NAVY
        p = doc.add_paragraph()
        r = p.add_run(a["customer"])
        r.font.size, r.font.color.rgb = Pt(18), BLUE
        for _ in range(3):
            doc.add_paragraph()
        ov = model["scores"].get("overall") or {}
        _p(doc, f"Overall health score: {ov.get('score', 'n/a')}/100   |   Risk: {ov.get('risk_level', 'n/a')}", bold=True, size=13, color=NAVY)
        _p(doc, f"QRadar version: {a.get('version') or model['environment'].get('version') or 'n/a'}", size=11)
        _p(doc, f"Prepared by: {a['consultant'] or 'IBM Security'}", size=11)
        _p(doc, f"Date: {a['generated']}", size=11)
        _p(doc, cfg["classification"].upper(), size=10, color=GREY)
        pb()
    # TOC
    doc.add_heading("Contents", 1)
    _add_field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
    _p(doc, "(Right-click and choose Update Field to refresh the table of contents after editing.)", italic=True, size=8, color=GREY)
    pb()

    S = model["scores"]
    cats, ov = S.get("categories") or {}, S.get("overall") or {}
    if "executive_summary" in secs:
        doc.add_heading("Executive Summary", 1)
        _p(doc, model["summary"].get("headline", ""), bold=True, size=12, color=BLUE)
        for t in model["summary"].get("paragraphs", []):
            _p(doc, t)
        _bullets(doc, model["summary"].get("highlights", []))
        if cfg.get("notes"):
            doc.add_heading("Consultant notes", 3)
            _comment(doc, cfg["notes"])
        sc("executive_summary")
        pb()
    if "health_score" in secs and ov:
        doc.add_heading("Health Score", 1)
        _img(doc, charts.gauge(ov["score"], f"Grade {ov['grade']}"), 8)
        _table(doc, [["Overall", "Risk score", "Risk level", "Maturity"], [f"{ov['score']}/100", f"{ov['risk_score']}/100", ov["risk_level"], f"{ov['maturity_level']} ({ov['maturity_index']}/5)"]])
        _img(doc, charts.category_bars(cats), 16)
        rows = [["Category", "Score", "Critical", "High", "Medium", "Low"]] + [[c["label"], c["score"], c["findings"]["critical"], c["findings"]["high"], c["findings"]["medium"], c["findings"]["low"]] for c in cats.values()]
        _table(doc, rows)
        _img(doc, charts.severity_donut(ov.get("severity_totals", {})), 7)
        sc("health_score")
        pb()
    if "environment" in secs:
        e = model["environment"]
        doc.add_heading("Environment Overview", 1)
        _table(doc, [["Attribute", "Value"], ["QRadar version", e.get("version") or "Not determined"], ["Fix pack", e.get("fix_pack") or "Not determined"], ["Build", e.get("build") or "Not determined"],
                     ["Deployment type", e.get("deployment_type") or "Not determined"], ["Installation type", e.get("install_type") or "Not determined"], ["Hosts", e.get("hosts")]], [5, 11])
        nodes = model["topology"].get("nodes", [])
        if nodes:
            doc.add_heading("Appliance inventory", 2)
            _table(doc, [["Hostname", "IP", "Role", "CPU", "RAM", "Disk"]] + [[n.get("hostname"), n.get("ip") or "-", n.get("role"), n.get("cpu_cores") or "-", _gb(n.get("memory_bytes")), _gb(n.get("disk_bytes"))] for n in nodes])
        sc("environment")
        pb()
    if "deployment" in secs:
        doc.add_heading("Deployment", 1)
        _img(doc, charts.deployment_diagram(model["topology"]), 16.5)
        doc.add_heading("Deployment validation", 2)
        _table(doc, [["Check", "Result", "Detail"]] + [[c["name"], c["status"], c["detail"]] for c in model["deployment_checks"]], [4.5, 2.5, 10])
        sc("deployment")
        pb()
    if "notifications" in secs:
        n = model["notifications"]
        doc.add_heading("Notification Analysis", 1)
        bs = n.get("by_severity", {})
        _table(doc, [["Critical", "High", "Medium", "Low", "Info", "Open"], [bs.get("critical", 0), bs.get("high", 0), bs.get("medium", 0), bs.get("low", 0), bs.get("info", 0), n.get("open", 0)]])
        for i in [x for x in n.get("items", []) if x["effective_severity"] in ("critical", "high", "medium")][:60]:
            doc.add_heading(f"{i['title']} ({i['effective_severity'].upper()} / {i['priority']})", 3)
            _p(doc, i["message"][:400], italic=True, size=9, color=GREY)
            _p(doc, "Root cause: " + i["root_cause"])
            _p(doc, "Business impact: " + i["business_impact"])
            _bullets(doc, i["recommendation"][:4])
            _p(doc, "IBM best practice: " + i["best_practice"], size=9, color=GREY)
            _comment(doc, i.get("comment"))
        sc("notifications")
        pb()
    if "logsources" in secs:
        l = model["logsources"]
        doc.add_heading("Log Source Analysis", 1)
        if l.get("total"):
            _table(doc, [["Total", "Enabled", "Disabled", "Error", "Unknown", "Total EPS", "Avg EPS"], [l["total"], l["enabled"], l["disabled"], l["error"], l["unknown"], f"{l['total_eps']:.0f}", f"{l['average_eps']:.1f}"]])
            pr = l.get("protocols", [])
            _img(doc, charts.bar([p["name"] for p in pr], [p["count"] for p in pr], "Log sources by protocol"), 14)
            _img(doc, charts.pie([p["name"] for p in pr], [p["eps"] for p in pr], "EPS by protocol"), 10)
            doc.add_heading("Top EPS sources", 2)
            _table(doc, [["Name", "Type", "Protocol", "EPS"]] + [[s["name"], s["type"], s["protocol"], f"{s['eps']:.1f}"] for s in l["top_eps"][:10]])
            if l.get("problematic"):
                doc.add_heading("Problematic log sources", 2)
                _table(doc, [["Name", "Type", "Status", "EPS", "Reason"]] + [[p["name"], p["type"], p["status"], f"{p['eps']:.1f}", "; ".join(p["reasons"])] for p in l["problematic"][:40]])
        else:
            _p(doc, "No log source data was present in the collected bundle.")
        for f in [f for f in model["findings"] if f["category"] == "logsources"][:8]:
            doc.add_heading(f["title"], 3)
            _p(doc, f["observation"])
            _bullets(doc, f["recommendation"][:4])
            _comment(doc, f.get("comment"))
        sc("logsources")
        pb()
    if "wincollect" in secs:
        w = model["wincollect"]
        doc.add_heading("WinCollect Analysis", 1)
        if w.get("total"):
            _table(doc, [["Total", "Managed", "Standalone", "Unmanaged", "Disconnected", "Outdated"], [w["total"], w["managed"], w["standalone"], w["unmanaged"], w["disconnected"], w["outdated"]]])
            v = w.get("versions", [])
            _img(doc, charts.bar([x["version"] for x in v][:12], [x["count"] for x in v][:12], "Agent version distribution", color=charts.TEAL, horizontal=False), 13)
            for title, key in (("Old versions", "old_versions"), ("Communication problems", "communication_problems"), ("Configuration issues", "config_issues")):
                if w.get(key):
                    doc.add_heading(title, 2)
                    _table(doc, [["Agent", "Version", "Status", "Detail"]] + [[x["name"], x.get("version") or "-", x["status"], x.get("reason") or ""] for x in w[key][:20]])
        else:
            _p(doc, "No WinCollect agent data was present in the collected bundle.")
        for f in [f for f in model["findings"] if f["category"] == "wincollect"][:5]:
            doc.add_heading(f["title"], 3)
            _p(doc, f["observation"])
            _bullets(doc, f["recommendation"][:3])
        sc("wincollect")
        pb()
    if "applications" in secs:
        ap = model["applications"]
        doc.add_heading("Application Analysis", 1)
        if ap.get("total"):
            _table(doc, [["Application", "Version", "Host", "Status", "CPU", "Memory", "Restarts"]] + [[i["name"], i.get("version") or "-", i.get("host") or "-", i["status"],
                    f"{i['cpu_pct']:.0f}%" if i.get("cpu_pct") is not None else "-", f"{i['memory_pct']:.0f}%" if i.get("memory_pct") is not None else "-", i.get("restart_count", 0)] for i in ap["items"][:50]])
            for i in [x for x in ap["items"] if x["issues"]][:15]:
                doc.add_heading(f"{i['name']} - {i['status']}", 3)
                _p(doc, "Observation: " + i["observation"])
                _p(doc, "Impact: " + i["impact"])
                _p(doc, "Priority: " + i["priority"])
                _bullets(doc, i["recommendation"][:3])
                _comment(doc, i.get("comment"))
        else:
            _p(doc, "No application data was present in the collected bundle.")
        sc("applications")
        pb()
    if "performance" in secs:
        doc.add_heading("Performance Analysis", 1)
        for name, h in list(model["performance_hosts"].items())[:6]:
            doc.add_heading(name, 2)
            ser = h.get("series", {})
            if ser.get("cpu"):
                _img(doc, charts.line(ser["cpu"], [("used", "CPU used %", charts.BLUE), ("iowait", "I/O wait %", "#ff832b")], "CPU utilisation", "%", (0, 100), hline=90), 16)
            if ser.get("memory"):
                _img(doc, charts.line(ser["memory"], [("used_pct", "Memory used %", charts.PURPLE)], "Memory utilisation", "%", (0, 100), hline=90), 16)
            if ser.get("disk"):
                _img(doc, charts.line(ser["disk"], [("await", "await (ms)", "#da1e28"), ("util", "%util", charts.TEAL)], "Disk latency & utilisation"), 16)
            if h.get("df"):
                _table(doc, [["Mount", "Size", "Used", "Use %"]] + [[f["mount"], _gb(f["size_bytes"]), _gb(f["used_bytes"]), f"{f['use_pct']:.0f}%"] for f in h["df"] if f.get("size_bytes")][:14])
        doc.add_heading("Performance findings", 2)
        rows = [["Host", "Metric", "Severity", "Finding", "Priority"]] + [[p["host"], p["label"], p["severity"], p["finding"], p["priority"]] for p in model["performance_assessments"] if p["severity"] not in ("ok", "info")][:60]
        if len(rows) > 1:
            _table(doc, rows, [2.5, 3, 2, 7.5, 1.5], sev_col=2)
        else:
            _p(doc, "No performance risks were identified.")
        sc("performance")
        pb()
    if "licensing" in secs:
        lic = model["licensing"]
        doc.add_heading("Licensing", 1)
        if lic.get("has_data"):
            rows = [["Metric", "Licensed", "Current", "Peak", "Consumption", "Risk"]]
            for m in (lic.get("eps"), lic.get("fpm")):
                if m:
                    rows.append([m["name"], f"{m['licensed']:,.0f}" if m.get("licensed") else "n/a", f"{m['current']:,.0f}" if m.get("current") is not None else "n/a",
                                 f"{m['peak']:,.0f}" if m.get("peak") is not None else "n/a", f"{m['consumption_pct']:.0f}%" if m.get("consumption_pct") is not None else "n/a", m["risk"]])
            _table(doc, rows)
            for m in (lic.get("eps"), lic.get("fpm")):
                if m and m.get("series"):
                    _img(doc, charts.licence_chart(m), 15)
        else:
            _p(doc, "No licence information was present in the collected bundle.")
        sc("licensing")
        pb()
    if "recommendations" in secs:
        doc.add_heading("Recommendations", 1)
        for t in model["top_recommendations"]:
            doc.add_heading(f"{t['rank']}. {t['title']} ({t['severity'].upper()} / {t['priority']})", 3)
            _p(doc, "Impact: " + t["impact"])
            _bullets(doc, t["recommendation"][:5])
            _p(doc, "IBM best practice: " + t["best_practice"], size=9, color=GREY)
        sc("recommendations")
        pb()
    if "next_steps" in secs:
        doc.add_heading("Next Steps", 1)
        for s in cfg["next_steps"]:
            doc.add_paragraph(s, style="List Number")
        sc("next_steps")
        pb()
    if "appendix" in secs:
        doc.add_heading("Appendix", 1)
        doc.add_heading("A. Complete findings register", 2)
        _table(doc, [["#", "Severity", "Priority", "Category", "Finding"]] + [[i, f["severity"], f["priority"], f["category"], f["title"]] for i, f in enumerate(model["findings"], 1)], [1, 2.2, 1.8, 3, 9], sev_col=1)
        _p(doc, f"Source package: {a['source_filename']}. Analysis is based exclusively on the data present in the supplied get_logs package.", size=9, color=GREY)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
